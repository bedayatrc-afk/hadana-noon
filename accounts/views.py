from django.shortcuts import render, redirect, get_object_or_404
from permissions_app.decorators import role_required
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.password_validation import validate_password
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponse
from django.contrib.auth.models import User
from django.contrib import messages
from django import forms
from django.db.models import Sum, Q
from django.core.mail import EmailMultiAlternatives
from django.conf import settings
from django.core.exceptions import ValidationError
from django.utils import timezone
from html import escape as html_escape

import openpyxl
import uuid

# ✅ النماذج الداخلية
from .forms import FeePaymentForm, EvaluationForm, ExpenseForm
from .models import Expense, FeePayment, UserProfile

# ✅ النماذج من accounts
from accounts.models import Student

# ✅ النماذج من parties
from parties.models import (
    ParentStudent,
    Person,
    StudentEvaluation,
    Attendance,
)

# ✅ النماذج المشتركة (محلية)
from .models import Evaluation, Parent


def _get_role(user):
    """Return the current user's role, repairing legacy accounts if needed."""
    fallback_role = "admin" if user.is_superuser else ("teacher" if user.is_staff else "student")
    profile, _ = UserProfile.objects.get_or_create(
        user=user, defaults={"role": fallback_role}
    )
    return profile.role


def _children_for_parent(user):
    """Support both the direct Student.parent link and the join-table link."""
    from parties.models import ParentStudent
    return Student.objects.filter(
        Q(parent=user) | Q(parents__parent=user)
    ).distinct()


def _can_access_student(user, student):
    role = _get_role(user)
    if role in {"admin", "teacher"}:
        return True
    if role == "parent":
        return _children_for_parent(user).filter(pk=student.pk).exists()
    if role == "student":
        return student.user_id == user.pk
    return False


def _evaluation_queryset_for_user(user):
    role = _get_role(user)
    qs = Evaluation.objects.select_related("student", "teacher")
    if role in {"admin", "teacher"}:
        return qs.all()
    if role == "parent":
        return qs.filter(student__in=_children_for_parent(user))
    if role == "student":
        return qs.filter(student__user=user)
    return qs.none()



# ──────────────────────────────────────────────
# 🏠 الصفحة الرئيسية
# ──────────────────────────────────────────────

@login_required
def home(request):
    return render(request, 'home.html')


# ──────────────────────────────────────────────
# 🔐 تسجيل الدخول والخروج
# ──────────────────────────────────────────────

def login_page(request):
    if request.user.is_authenticated:
        _get_role(request.user)
        return redirect("main_dashboard")
    if request.method == "POST":
        username = request.POST.get("username")
        password = request.POST.get("password")
        user = authenticate(request, username=username, password=password)
        if user is not None:
            _get_role(user)
            login(request, user)
            return redirect("main_dashboard")
        else:
            messages.error(request, "⚠️ المستخدم غير موجود أو كلمة المرور غير صحيحة")
    return render(request, "login.html")


def logout_user(request):
    logout(request)
    return redirect("login")


def landing_page(request):
    if request.user.is_authenticated:
        _get_role(request.user)
        return redirect("main_dashboard")
    return render(request, "login.html")


@login_required
@role_required("admin", "teacher", "parent", "student")
def student_code_search(request):
    student_info = None
    error = None
    
    if request.method == "POST":
        student_code = (request.POST.get("student_code") or "").strip().upper()
        if not student_code:
            error = "اكتب كود الطالب أولًا."
        else:
            student = Student.objects.filter(student_code=student_code).first()
            if student is None:
                error = "كود الطالب غير صحيح"
            elif _can_access_student(request.user, student):
                student_info = student
            else:
                error = "لا تملك صلاحية عرض بيانات هذا الطالب."
    
    return render(request, "code.html", {"student_info": student_info, "error": error})


# ──────────────────────────────────────────────
# 📊 الداشبوردات
# ──────────────────────────────────────────────

@login_required
@role_required("admin", "teacher", "parent", "student")
def main_dashboard(request):
    return render(request, "main_dashboard.html")


@login_required
def dashboard(request):
    return render(request, "dashboard.html")


@login_required
def teacher_dashboard(request):
    return render(request, "teacher_dashboard.html")


@login_required
@role_required("parent")
def parent_dashboard(request):
    students = _children_for_parent(request.user)

    if not students.exists():
        return redirect('link_student')

    student = students.first()

    evaluations = Evaluation.objects.filter(student__in=students).order_by('-date')
    attendances = Attendance.objects.filter(student__in=students).order_by('-date')
    payments = FeePayment.objects.filter(student__in=students).order_by('-payment_date')

    total_absences = attendances.filter(status='absent').count()
    total_evaluations = evaluations.count()
    total_paid = payments.filter(status='confirmed').aggregate(total=Sum('amount'))['total'] or 0

    # ✅ أسبوع ابنك بالصور: بناء آخر 7 أيام من بيانات حقيقية
    from datetime import date, timedelta
    day_names_ar = {5: "سبت", 6: "أحد", 0: "اتنين", 1: "ثلاثاء", 2: "أربعاء", 3: "خميس", 4: "جمعة"}

    week_data = []
    today = timezone.localdate()
    for i in range(6, -1, -1):
        d = today - timedelta(days=i)
        eval_today = evaluations.filter(date=d).first()
        attendance_today = attendances.filter(date=d).first()

        if eval_today:
            mood = eval_today.mood  # happy / sad / tired / sick / sleeping
        elif attendance_today and attendance_today.status == 'absent':
            mood = 'home'
        else:
            mood = 'none'

        week_data.append({
            'day_name': day_names_ar[d.weekday()],
            'mood': mood,
        })

    context = {
        'student': student,
        'evaluations': evaluations[:20],
        'attendances': attendances[:20],
        'payments': payments[:20],
        'total_absences': total_absences,
        'total_evaluations': total_evaluations,
        'total_paid': total_paid,
        'week_data': week_data,
    }
    return render(request, 'accounts/parent_dashboard.html', context)


# ──────────────────────────────────────────────
# 📝 التسجيل
# ──────────────────────────────────────────────

@login_required
@role_required("admin", "teacher")
def register_student(request):
    if request.method == "POST":
        name = (request.POST.get("name") or "").strip()
        phone = (request.POST.get("phone") or "").strip()
        grade_raw = (request.POST.get("grade") or "0").strip()
        if not name:
            messages.error(request, "اكتب اسم الطالب أولًا.")
            return redirect("register_student")
        try:
            grade = max(0, int(grade_raw or 0))
        except ValueError:
            messages.error(request, "الصف الدراسي لازم يكون رقمًا صحيحًا.")
            return redirect("register_student")

        user = User.objects.create(username=f"student_{uuid.uuid4().hex[:8]}")
        user.set_unusable_password()
        user.save(update_fields=["password"])

        student = Student.objects.create(
            user=user, name=name, phone=phone, grade=grade
        )
        Person.objects.update_or_create(
            user=user,
            defaults={
                "name": student.name,
                "role": "student",
                "phone": phone or "",
                "national_id": "",
                "email": user.email or "",
            },
        )

        messages.success(
            request,
            f"✅ تم تسجيل الطالب بنجاح - كود الطالب هو: {student.student_code}"
        )
        return redirect("register_student")

    return render(request, "register_student.html")


@login_required
@role_required("admin", "teacher")
def register_parent(request):
    if request.method == "POST":
        parent_name = (request.POST.get("parent_name") or "").strip()
        phone = (request.POST.get("phone") or "").strip()
        if not parent_name or not phone:
            messages.error(request, "اكتب اسم ولي الأمر ورقم الهاتف.")
            return redirect("register_parent")

        Parent.objects.create(parent_name=parent_name, phone=phone)
        messages.success(request, "✅ تم التسجيل بنجاح")
        return redirect("register_parent")

    return render(request, "register_parent.html")


@login_required
@role_required("admin", "teacher")
def register_user(request):
    if request.method == "POST":
        username = (request.POST.get("username") or "").strip()
        email = (request.POST.get("email") or "").strip()
        password = request.POST.get("password") or ""
        role = (request.POST.get("role") or "teacher").strip()
        allowed_roles = {key for key, _label in UserProfile.ROLE_CHOICES if key != "admin"}

        if not username or not password:
            messages.error(request, "اسم المستخدم وكلمة المرور مطلوبان.")
            return redirect("register_user")
        if role not in allowed_roles:
            messages.error(request, "الدور المحدد غير صالح.")
            return redirect("register_user")
        if User.objects.filter(username__iexact=username).exists():
            messages.error(request, "اسم المستخدم موجود بالفعل ❌")
            return redirect("register_user")
        if len(password) < 8:
            messages.error(request, "كلمة المرور يجب ألا تقل عن 8 أحرف.")
            return redirect("register_user")
        try:
            validate_password(password, User(username=username, email=email))
        except ValidationError as exc:
            messages.error(request, " ".join(exc.messages))
            return redirect("register_user")

        user = User.objects.create_user(username=username, email=email, password=password)
        UserProfile.objects.update_or_create(user=user, defaults={"role": role})
        messages.success(request, "تم إنشاء الحساب وتحديد صلاحياته بنجاح ✅")
        return redirect("register_user")

    return render(request, "register_user.html", {"roles": [item for item in UserProfile.ROLE_CHOICES if item[0] != "admin"]})


# ──────────────────────────────────────────────
# 🔗 ربط الطالب بولي الأمر
# ──────────────────────────────────────────────

@login_required
@role_required("parent")
def link_student(request):
    if request.method == 'POST':
        student_code = (request.POST.get('student_code') or '').strip().upper()
        try:
            student = Student.objects.get(student_code=student_code)
            if student.parent_id and student.parent_id != request.user.id:
                messages.error(request, "هذا الطالب مرتبط بالفعل بحساب ولي أمر آخر.")
                return redirect('link_student')
            student.parent = request.user
            student.save(update_fields=["parent"])
            from parties.models import ParentStudent
            ParentStudent.objects.get_or_create(parent=request.user, student=student)
            messages.success(request, f"✅ تم ربط الطالب '{student.name}' بنجاح!")
            return redirect('parent_dashboard')
        except Student.DoesNotExist:
            messages.error(request, "❌ كود الطالب غير صحيح.")

    return render(request, 'link_student.html')


# ──────────────────────────────────────────────
# 💵 المدفوعات
# ──────────────────────────────────────────────

@login_required
@role_required("admin")
def pay_fees(request):
    if request.method == 'POST':
        form = FeePaymentForm(request.POST, request.FILES)
        if form.is_valid():
            payment = form.save(commit=False)
            payment.user = request.user
            payment.student = form.cleaned_data['student']
            payment.save()
            messages.success(request, "تم تسجيل الدفعة بنجاح.")
            return redirect("payments_list")
    else:
        form = FeePaymentForm()
    return render(request, 'pay_fees.html', {'form': form})


@login_required
def payments_list(request):
    role = _get_role(request.user)
    payments = FeePayment.objects.select_related("student", "student__user", "user")
    if role == "parent":
        payments = payments.filter(student__in=_children_for_parent(request.user))
    elif role == "student":
        payments = payments.filter(student__user=request.user)
    elif role != "admin":
        payments = payments.none()
    total_amount = payments.aggregate(total=Sum('amount'))['total'] or 0
    return render(request, "payments_list.html", {
        "payments": payments.order_by("-payment_date"),
        "total_amount": total_amount
    })


@login_required
def print_payment(request, id):
    payment = get_object_or_404(FeePayment.objects.select_related("student", "student__user", "user"), id=id)
    role = _get_role(request.user)
    if role not in {"admin", "parent", "student"} or (role != "admin" and not _can_access_student(request.user, payment.student)):
        from django.http import Http404
        raise Http404
    return render(request, "print_payment.html", {"payment": payment})


@login_required
def add_payment(request):
    return render(request, 'add_payment.html')


@login_required
@role_required("admin")
def export_payments_excel(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Payments"

    ws.append([
        "👦 الطالب", "💵 المبلغ", "🕒 تاريخ الدفع", "💳 طريقة الدفع",
        "🔢 رقم العملية", "📱 الهاتف", "📌 الحالة", "📝 ملاحظات", "👤 المستخدم"
    ])

    role = _get_role(request.user)
    payments = FeePayment.objects.select_related("student", "user")
    if role == "parent":
        payments = payments.filter(student__in=_children_for_parent(request.user))
    elif role == "student":
        payments = payments.filter(student__user=request.user)
    elif role != "admin":
        payments = payments.none()
    for p in payments.order_by("-payment_date"):

        student_display = (p.student.name or p.student.user.username) if p.student and p.student.user else (p.student.name if p.student else "---")
        user_username = p.user.username if p.user else "---"

        ws.append([
            student_display,
            p.amount,
            p.payment_date.strftime("%d/%m/%Y %H:%M") if p.payment_date else "---",
            p.get_payment_method_display() if p.payment_method else "---",
            p.transaction_id or "---",
            p.sender_phone or "---",
            p.get_status_display() if p.status else "---",
            p.notes or "---",
            user_username
        ])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response['Content-Disposition'] = 'attachment; filename=payments.xlsx'
    wb.save(response)
    return response


# ──────────────────────────────────────────────
# 📋 التقييمات
# ──────────────────────────────────────────────

@login_required
@role_required("admin", "teacher")
def evaluation(request):
    if request.method == "POST":
        form = EvaluationForm(request.POST)
        if form.is_valid():
            eval_obj = form.save(commit=False)
            eval_obj.teacher = request.user

            student_id = request.POST.get("student")
            if student_id:
                try:
                    student = Student.objects.get(id=student_id)
                    eval_obj.student = student
                except Student.DoesNotExist:
                    messages.error(request, "الطالب غير موجود ❌")
                    return redirect("evaluation")

            eval_obj.save()
            messages.success(request, "تم تسجيل التقييم بنجاح ✅")
            return redirect("evaluations_list")
    else:
        form = EvaluationForm()

    students = Student.objects.select_related("user").all()
    return render(request, "add_evaluation.html", {
        "form": form,
        "students": students
    })


@login_required
def evaluations_list(request):
    evaluations = _evaluation_queryset_for_user(request.user).order_by("-date", "-id")
    return render(request, "evaluations_list.html", {"evaluations": evaluations})


@login_required
def print_evaluation(request, id):
    evaluation = get_object_or_404(Evaluation.objects.select_related("student", "teacher"), id=id)
    if not _can_access_student(request.user, evaluation.student):
        from django.http import Http404
        raise Http404
    return render(request, "print_evaluation.html", {"evaluation": evaluation})


# ✅ فورم إدخال إيميل ولي الأمر
class GuardianEmailForm(forms.Form):
    guardian_email = forms.EmailField(label="إيميل ولي الأمر")


@login_required
def send_evaluation_mail(request, id):
    evaluation = get_object_or_404(Evaluation.objects.select_related("student", "teacher"), id=id)
    if not _can_access_student(request.user, evaluation.student):
        from django.http import Http404
        raise Http404

    if request.method == "POST":
        form = GuardianEmailForm(request.POST)
        if form.is_valid():
            recipient_email = form.cleaned_data["guardian_email"]
            if not settings.EMAIL_HOST_USER or not settings.EMAIL_HOST_PASSWORD:
                messages.error(
                    request,
                    "لم يُرسل البريد فعليًا لأن إعدادات SMTP غير مكتملة. أضف بيانات البريد إلى ملف .env ثم أعد تشغيل الخادم."
                )
                return render(request, "send_mail_form.html", {"form": form, "evaluation": evaluation})

            subject = f"تقرير الطالب {evaluation.student.name}"
            text_content = f"ده تقرير الطالب {evaluation.student.name} بتاريخ {evaluation.date.strftime('%d/%m/%Y')}"

            html_content = f"""
            <html>
              <body>
                <h2>📋 تقرير الطالب</h2>
                <p>👦 الطالب: {html_escape(evaluation.student.name)}</p>
                <p>📅 التاريخ: {evaluation.date.strftime("%d/%m/%Y")}</p>
                <p>🙂 الحالة: {html_escape(evaluation.get_mood_display())}</p>
                <p>📚 المادة: {html_escape(evaluation.get_subject_display() or "غير محدد")}</p>
                <p>📈 الاستجابة: {html_escape(evaluation.get_response_display())}</p>
                <p>🧑‍🍼 السلوك: {html_escape(evaluation.get_behavior_display())}</p>
                <p>🚻 الحمام: من {html_escape(str(evaluation.bathroom_from or "—"))} إلى {html_escape(str(evaluation.bathroom_to or "—"))}</p>
                <p>😴 النوم: من {html_escape(str(evaluation.sleep_from or "—"))} إلى {html_escape(str(evaluation.sleep_to or "—"))}</p>
                <p>📝 الملاحظات: {html_escape(evaluation.notes or "—")}</p>
                <p>👩‍🏫 المعلم: {html_escape(evaluation.teacher.username)}</p>
              </body>
            </html>
            """

            msg = EmailMultiAlternatives(
                subject, text_content, settings.DEFAULT_FROM_EMAIL, [recipient_email]
            )
            msg.attach_alternative(html_content, "text/html")
            try:
                msg.send(fail_silently=False)
            except Exception:
                messages.error(request, "تعذّر إرسال البريد. راجع إعدادات البريد في ملف .env وسجل الأخطاء.")
            else:
                messages.success(
                    request,
                    f"📧 تم إرسال التقييم للطالب {evaluation.student.name} إلى {recipient_email}"
                )
                return redirect("evaluations_list")
    else:
        form = GuardianEmailForm()

    return render(request, "send_mail_form.html", {"form": form, "evaluation": evaluation})


@login_required
@role_required("admin", "teacher")
def export_evaluations_excel(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Evaluations"

    ws.append([
        "👦 الطالب", "📅 التاريخ", "🙂 الحالة", "📚 المادة",
        "📈 الاستجابة", "🧑‍🍼 السلوك", "🚻 الحمام", "😴 النوم",
        "📝 الملاحظات", "👩‍🏫 المعلمة"
    ])

    for e in Evaluation.objects.select_related("student", "teacher").all():
        ws.append([
            e.student.name,
            e.date.strftime("%d/%m/%Y"),
            e.get_mood_display(),
            e.get_subject_display(),
            e.get_response_display(),
            e.get_behavior_display(),
            f"{e.bathroom_from or ''} - {e.bathroom_to or ''}",
            f"{e.sleep_from or ''} - {e.sleep_to or ''}",
            e.notes or "---",
            e.teacher.username
        ])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response['Content-Disposition'] = 'attachment; filename=evaluations.xlsx'
    wb.save(response)
    return response


@login_required
@role_required("admin", "teacher")
def export_evaluations(request):
    return export_evaluations_excel(request)


# ──────────────────────────────────────────────
# 💰 المصاريف
# ──────────────────────────────────────────────

@login_required
@role_required("admin")
def expenses(request):
    if request.method == "POST":
        form = ExpenseForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "تم تسجيل المصروف بنجاح.")
            return redirect("expenses")
    else:
        form = ExpenseForm()
    expenses_list = Expense.objects.all().order_by('-date')
    total_expenses = Expense.objects.aggregate(total=Sum('amount'))['total'] or 0
    return render(request, "expenses.html", {
        "expenses": expenses_list,
        "total_expenses": total_expenses,
        "form": form,
    })


# ──────────────────────────────────────────────
# 📊 التقارير والإعدادات
# ──────────────────────────────────────────────

@login_required
@role_required("admin")
def reports(request):
    evaluations = Evaluation.objects.all().order_by('-date')
    return render(request, 'reports.html', {
        'evaluations': evaluations,
        'students_count': Student.objects.count(),
        'evaluations_count': Evaluation.objects.count(),
        'payments_total': FeePayment.objects.aggregate(Sum('amount'))['amount__sum'] or 0,
    })


@login_required
@role_required("admin")
def settings_page(request):
    if request.method == "POST":
        user = request.user
        username = (request.POST.get("username") or user.username).strip()
        email = (request.POST.get("email") or user.email).strip()
        if not username:
            messages.error(request, "اسم المستخدم لا يمكن أن يكون فارغًا.")
            return redirect("settings")
        if User.objects.filter(username__iexact=username).exclude(pk=user.pk).exists():
            messages.error(request, "اسم المستخدم مستخدم بالفعل.")
            return redirect("settings")
        user.username = username
        user.email = email
        user.first_name = (request.POST.get("first_name") or "").strip()
        user.last_name = (request.POST.get("last_name") or "").strip()
        user.save()
        messages.success(request, "تم حفظ التغييرات بنجاح ✅")
        return redirect("settings")

    return render(request, "settings.html")


# ──────────────────────────────────────────────
# 🖊️ إنشاء حساب (Register)
# ──────────────────────────────────────────────

class RegisterForm(forms.Form):
    username = forms.CharField(
        label="اسم المستخدم",
        max_length=150,
        widget=forms.TextInput(attrs={"autocomplete": "username", "placeholder": "اسم المستخدم"}),
    )
    email = forms.EmailField(
        label="البريد الإلكتروني",
        widget=forms.EmailInput(attrs={"autocomplete": "email", "placeholder": "name@example.com"}),
    )
    password = forms.CharField(
        label="كلمة المرور",
        widget=forms.PasswordInput(attrs={"autocomplete": "new-password"}),
    )
    confirm_password = forms.CharField(
        label="تأكيد كلمة المرور",
        widget=forms.PasswordInput(attrs={"autocomplete": "new-password"}),
    )

    def clean_username(self):
        username = self.cleaned_data["username"].strip()
        if User.objects.filter(username__iexact=username).exists():
            raise forms.ValidationError("اسم المستخدم مستخدم بالفعل.")
        return username

    def clean(self):
        cleaned = super().clean()
        password = cleaned.get("password")
        confirm_password = cleaned.get("confirm_password")
        if password and confirm_password and password != confirm_password:
            self.add_error("confirm_password", "كلمتا المرور غير متطابقتين.")
        if password:
            try:
                validate_password(password, User(username=cleaned.get("username", ""), email=cleaned.get("email", "")))
            except ValidationError as exc:
                self.add_error("password", exc)
        return cleaned


def register_page(request):
    if request.method == "POST":
        form = RegisterForm(request.POST)
        if form.is_valid():
            user = User.objects.create_user(
                username=form.cleaned_data["username"],
                email=form.cleaned_data["email"],
                password=form.cleaned_data["password"]
            )
            Student.objects.create(user=user, name=user.username, phone="", grade=0)
            UserProfile.objects.update_or_create(user=user, defaults={"role": "student"})
            login(request, user)
            messages.success(request, "تم إنشاء حسابك بنجاح. يمكنك استكمال بياناتك من لوحة التحكم.")
            return redirect("main_dashboard")
        else:
            return render(request, "register.html", {"form": form})
    else:
        form = RegisterForm()
    return render(request, "register.html", {"form": form})


# ──────────────────────────────────────────────
# 🔍 Autocomplete والحضور
# ──────────────────────────────────────────────

@login_required
@role_required("admin", "teacher")
def student_autocomplete(request):
    q = (request.GET.get("q") or "").strip()
    if not q:
        return JsonResponse([], safe=False)
    results = list(Student.objects.filter(
        Q(name__icontains=q) | Q(user__username__icontains=q)
    ).values_list("name", flat=True).distinct()[:20])
    return JsonResponse(results, safe=False)


class AttendanceForm(forms.ModelForm):
    class Meta:
        model = Attendance
        fields = ['student', 'date', 'status']
        widgets = {
            'date': forms.DateInput(attrs={'type': 'date'}),
        }
@login_required
@role_required("admin", "teacher")
def add_attendance(request):
    form = AttendanceForm()
    if request.method == 'POST':
        form = AttendanceForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "✅ تم تسجيل الحضور بنجاح")
            return redirect('add_attendance')
    return render(request, 'add_attendance.html', {'form': form})

@login_required
@role_required("admin", "teacher")
def attendance_list(request):
    attendances = Attendance.objects.select_related(
        "student", "student__user"
    ).all().order_by('-date')

    present_count = 0
    absent_count = 0
    for a in attendances:
        status_val = a.status or ""
        if status_val in ('present', 'حاضر', 'حضور'):
            present_count += 1
        elif status_val in ('absent', 'غياب', 'غائب', 'غائبة'):
            absent_count += 1

    return render(request, 'attendance_list.html', {
        'attendances': attendances,
        'present_count': present_count,
        'absent_count': absent_count,
    })


# ──────────────────────────────────────────────
# 📊 تصدير الحضور لـ Excel
# ──────────────────────────────────────────────

@login_required
@role_required("admin", "teacher")
def export_attendance_excel(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance"

    ws.append(["👦 الطالب", "📅 التاريخ", "📌 الحالة"])

    for a in Attendance.objects.select_related("student", "student__user").all().order_by('-date'):
        if a.student and a.student.name:
            student_name = a.student.name
        elif a.student and a.student.user:
            student_name = a.student.user.username
        else:
            student_name = "---"

        status_display = a.get_status_display() if hasattr(a, 'get_status_display') else (a.status or "---")

        ws.append([
            student_name,
            a.date.strftime("%d/%m/%Y") if a.date else "---",
            status_display,
        ])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response['Content-Disposition'] = 'attachment; filename=attendance.xlsx'
    wb.save(response)
    return response
